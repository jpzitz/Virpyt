#VirPyT test-03
#trying to jsut make everything work first before putting in classes


import openpyxl
#from openpyxl import Workbook


wb = openpyxl.load_workbook(filename = input("Input filename: "))

worksheet = wb.active
if worksheet:
    print("ok!")

for sheet in wb.sheetnames:
    print("Found sheet named %s" %sheet.title())

    
    
    #for table in sheet.tables:
     #   print("Found table named %t" %table.title)

     #   for column_header in table.header:  # iterate through cells that make
                                            # up the header
      #     print(“col header is %s” % column_header.value)
       # for row in table:
        #    cell = row['<col name>']        # we want to be able to access a
                                            # cell according to the header name.
         #   print('row info is %s' % cell.value)
